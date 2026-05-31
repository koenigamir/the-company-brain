"""Shared ingestion utilities: extract, chunk, and append documents to the brain."""

from __future__ import annotations

import os
import random
import datetime
import tempfile
import fitz  # PyMuPDF
from docx import Document as DocxDocument
from openpyxl import load_workbook

try:
    from langchain.text_splitter import RecursiveCharacterTextSplitter
except ModuleNotFoundError:
    from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_chroma import Chroma

import graph_engine

DATA_DIR = os.getenv("COMPANY_BRAIN_DATA_DIR", "data")
CHROMA_DIR = os.getenv("COMPANY_BRAIN_CHROMA_DIR", "chroma_db")
COLLECTION_NAME = "company_brain"
EMBED_MODEL = "sentence-transformers/all-MiniLM-L6-v2"
CHUNK_SIZE = 500
CHUNK_OVERLAP = 50

ROLE_OWNERS = [
    "ESG Compliance",
    "Master Data Ops",
    "Tax Team",
    "Regulatory Services",
    "Product Coverage & Onboarding",
    "Compliance & Sanctions",
]

ROLE_KEYWORDS = {
    "ESG Compliance": ["esg", "sfdr", "taxonomy"],
    "Tax Team": ["fatca", "tax"],
    "Master Data Ops": ["master", "mifid", "mifir", "data", "emt", "attribute"],
    "Regulatory Services": ["regulatory navigator", "product governance", "suitability", "complexity"],
    "Product Coverage & Onboarding": ["product coverage", "coverage", "onboarding"],
    "Compliance & Sanctions": ["aml", "kyc", "sanctions", "surveillance"],
}


def assign_role_owner(filename: str) -> str:
    name = filename.lower()
    for role, keywords in ROLE_KEYWORDS.items():
        if any(k in name for k in keywords):
            return role
    return random.choice(ROLE_OWNERS)


def today_iso() -> str:
    return datetime.date.today().isoformat()


def file_modified_iso(path: str) -> str:
    """Return the file's last-modified date (ISO). Falls back to today."""
    try:
        ts = os.path.getmtime(path)
        return datetime.date.fromtimestamp(ts).isoformat()
    except OSError:
        return today_iso()


def read_pdf(path: str) -> str:
    text_parts = []
    with fitz.open(path) as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts)


def read_docx(path: str) -> str:
    doc = DocxDocument(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def read_xlsx(path: str) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows_text = []
    for ws in wb.worksheets:
        rows_text.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows_text.append(" | ".join(cells))
    wb.close()
    return "\n".join(rows_text)


def read_txt(path: str) -> str:
    with open(path, encoding="utf-8", errors="ignore") as f:
        return f.read()


def extract_text(path: str) -> str:
    try:
        text = read_pdf(path)
        if text.strip():
            return text
    except Exception:
        pass

    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".txt":
            return read_txt(path)
        if ext == ".docx":
            return read_docx(path)
        if ext in (".xlsx", ".xlsm"):
            return read_xlsx(path)
    except Exception:
        return ""
    return ""


def get_splitter() -> RecursiveCharacterTextSplitter:
    return RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP
    )


def file_to_documents(
    path: str,
    filename: str | None = None,
    role_owner: str | None = None,
    last_updated: str | None = None,
) -> list[Document]:
    """Parse a file on disk into chunked LangChain Documents."""
    filename = filename or os.path.basename(path)
    text = extract_text(path)
    if not text.strip():
        return []

    role_owner = role_owner or assign_role_owner(filename)
    last_updated = last_updated or file_modified_iso(path)
    splitter = get_splitter()

    return [
        Document(
            page_content=chunk,
            metadata={
                "source_file": filename,
                "role_owner": role_owner,
                "last_updated": last_updated,
            },
        )
        for chunk in splitter.split_text(text)
    ]


def text_to_documents(
    text: str,
    filename: str,
    role_owner: str,
    last_updated: str,
    modality: str = "document",
    extractor: str = "",
    confidence: float = 0.0,
) -> list[Document]:
    """Chunk already-extracted text into Documents, carrying modality metadata.

    Used for inputs where text was produced by an extractor (e.g. image OCR)
    rather than read directly from the file on disk.
    """
    if not text.strip():
        return []
    meta = {
        "source_file": filename,
        "role_owner": role_owner,
        "last_updated": last_updated,
        "source_modality": modality,
    }
    if extractor:
        meta["extractor"] = extractor
        meta["extraction_confidence"] = round(float(confidence), 3)
    # Keep image OCR as one chunk when small — tables stay intact for retrieval.
    if modality == "image" and len(text) <= 4000:
        return [Document(page_content=text.strip(), metadata=dict(meta))]
    return [
        Document(page_content=chunk, metadata=dict(meta))
        for chunk in get_splitter().split_text(text)
    ]


def _fmt_ts(seconds: float) -> str:
    seconds = int(seconds)
    return f"{seconds // 60:02d}:{seconds % 60:02d}"


def segments_to_documents(
    segments: list[dict],
    filename: str,
    role_owner: str,
    last_updated: str,
    modality: str = "audio",
    extractor: str = "",
    confidence: float = 0.0,
) -> list[Document]:
    """Group transcript segments into ~CHUNK_SIZE chunks carrying time refs."""
    docs: list[Document] = []
    buf: list[str] = []
    buf_len = 0
    start = segments[0]["start"] if segments else 0.0
    end = start

    def flush(seg_start: float, seg_end: float):
        if not buf:
            return
        docs.append(
            Document(
                page_content=" ".join(buf).strip(),
                metadata={
                    "source_file": filename,
                    "role_owner": role_owner,
                    "last_updated": last_updated,
                    "source_modality": modality,
                    "time_ref": f"{_fmt_ts(seg_start)}-{_fmt_ts(seg_end)}",
                    "extractor": extractor,
                    "extraction_confidence": round(float(confidence), 3),
                },
            )
        )

    for seg in segments:
        text = (seg.get("text") or "").strip()
        if not text:
            continue
        if not buf:
            start = seg["start"]
        buf.append(text)
        buf_len += len(text) + 1
        end = seg["end"]
        if buf_len >= CHUNK_SIZE:
            flush(start, end)
            buf, buf_len = [], 0
    flush(start, end)
    return docs


def get_chroma_store() -> Chroma:
    embeddings = HuggingFaceEmbeddings(model_name=EMBED_MODEL)
    return Chroma(
        collection_name=COLLECTION_NAME,
        embedding_function=embeddings,
        persist_directory=CHROMA_DIR,
    )


def apply_access_metadata(
    docs: list[Document],
    visibility_roles: list[str],
    min_clearance: str,
) -> list[Document]:
    visibility_value = ", ".join(visibility_roles)
    for doc in docs:
        doc.metadata["visibility_roles"] = visibility_value
        doc.metadata["min_clearance"] = min_clearance
    return docs


def add_file_to_brain(
    file_bytes: bytes,
    filename: str,
    role_owner: str | None = None,
    visibility_roles=None,
    min_clearance: str | None = None,
    save_to_data_dir: bool = True,
) -> dict:
    """
    Ingest an uploaded file: save to data/, chunk, embed into Chroma, update graph.

    Returns a summary dict for the UI.
    """
    os.makedirs(DATA_DIR, exist_ok=True)

    dest_path = os.path.join(DATA_DIR, filename)
    if save_to_data_dir:
        with open(dest_path, "wb") as f:
            f.write(file_bytes)
        parse_path = dest_path
    else:
        with tempfile.NamedTemporaryFile(
            delete=False, suffix=os.path.splitext(filename)[1]
        ) as tmp:
            tmp.write(file_bytes)
            parse_path = tmp.name

    import media_ingest
    import image_ingest

    modality = media_ingest.modality_of(filename)
    extractor = ""
    confidence = 0.0
    is_image_doc = False

    try:
        if media_ingest.is_media(filename):
            transcript = media_ingest.transcribe(parse_path, filename=filename)
            text = transcript["text"]
            extractor = transcript.get("extractor", "")
            confidence = transcript.get("confidence", 0.0)
            segments = transcript.get("segments") or []
        elif image_ingest.is_image(filename):
            extracted = image_ingest.extract(parse_path, filename=filename)
            text = extracted["text"]
            extractor = extracted.get("extractor", "")
            confidence = extracted.get("confidence", 0.0)
            modality = "image"
            is_image_doc = True
            segments = None
        else:
            text = extract_text(parse_path)
            segments = None

        if not text.strip():
            return {
                "ok": False,
                "filename": filename,
                "error": "Could not extract text/transcript from this file.",
            }

        # Dynamic role assignment: match against the SIX role catalog. A
        # document may be owned by several roles when it spans domains.
        role_reason = ""
        role_owners = [role_owner] if role_owner else []
        if role_owner is None:
            from role_resolver import resolve_role

            resolved = resolve_role(text, filename)
            role_owner = resolved["role"]
            role_owners = resolved.get("roles") or [role_owner]
            role_reason = resolved["reason"]

        import store

        access = store.normalize_document_access(
            source_file=filename,
            role_owner=role_owner,
            role_owners=role_owners,
            visibility_roles=visibility_roles,
            min_clearance=min_clearance,
        )
        role_owners = access["role_owners"] or [role_owner]
        visibility_roles = access["visibility_roles"]
        min_clearance = access["min_clearance"]

        last_updated = file_modified_iso(parse_path)
        if segments is not None:
            docs = segments_to_documents(
                segments,
                filename=filename,
                role_owner=role_owner,
                last_updated=last_updated,
                modality=modality,
                extractor=extractor,
                confidence=confidence,
            )
            # Fallback: transcript text with no usable segments still gets chunked.
            if not docs and text.strip():
                docs = [
                    Document(
                        page_content=chunk,
                        metadata={
                            "source_file": filename,
                            "role_owner": role_owner,
                            "last_updated": last_updated,
                            "source_modality": modality,
                            "extractor": extractor,
                            "extraction_confidence": round(float(confidence), 3),
                        },
                    )
                    for chunk in get_splitter().split_text(text)
                ]
        elif is_image_doc:
            docs = text_to_documents(
                text,
                filename=filename,
                role_owner=role_owner,
                last_updated=last_updated,
                modality="image",
                extractor=extractor,
                confidence=confidence,
            )
        else:
            docs = file_to_documents(
                parse_path,
                filename=filename,
                role_owner=role_owner,
                last_updated=last_updated,
            )

        # Record all owning roles on each chunk (Chroma needs a scalar value).
        if not role_owners:
            role_owners = [role_owner]
        owners_str = ", ".join(role_owners)
        for d in docs:
            d.metadata["role_owners"] = owners_str
        apply_access_metadata(docs, visibility_roles, min_clearance)
    finally:
        if not save_to_data_dir:
            os.unlink(parse_path)

    if not docs:
        return {
            "ok": False,
            "filename": filename,
            "error": "Could not extract text from this file.",
        }

    vstore = get_chroma_store()
    vstore.add_documents(docs)

    graph = graph_engine.load_graph() or {
        "documents": {},
        "entities": {},
        "roles": {},
    }
    chunk_texts = [d.page_content for d in docs]
    role = docs[0].metadata["role_owner"]
    graph = graph_engine.merge_document(graph, filename, role, chunk_texts)
    graph_engine.save_graph(graph)

    # Persist the document ownership record (Supabase or local fallback).
    try:
        store.upsert_document(
            filename,
            role,
            docs[0].metadata["last_updated"],
            len(docs),
            modality=modality,
            role_owners=role_owners,
            visibility_roles=visibility_roles,
            min_clearance=min_clearance,
        )
    except Exception:
        pass

    return {
        "ok": True,
        "filename": filename,
        "chunks": len(docs),
        "role_owner": role,
        "role_owners": role_owners,
        "role_reason": role_reason,
        "modality": modality,
        "extractor": extractor,
        "extraction_confidence": round(float(confidence), 3) if extractor else None,
        "entities": graph["documents"].get(filename, {}).get("entities", []),
    }
