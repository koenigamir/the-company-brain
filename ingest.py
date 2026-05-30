import os
import shutil
import random
import datetime
import fitz  # PyMuPDF
from docx import Document as DocxDocument
from openpyxl import load_workbook
from dotenv import load_dotenv

from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_chroma import Chroma

load_dotenv()

DATA_DIR = "data"
CHROMA_DIR = "chroma_db"
COLLECTION_NAME = "company_brain"
EMBED_MODEL = "sentence-transformers/all-MiniLM-L6-v2"

ROLE_OWNERS = ["ESG Compliance", "Master Data Ops", "Tax Team"]

# A smarter-than-random mapping for a convincing demo. Falls back to random
# if no keyword matches. (Spec allows pure random; this just looks better.)
ROLE_KEYWORDS = {
    "ESG Compliance": ["esg", "sfdr", "taxonomy"],
    "Tax Team": ["fatca", "tax"],
    "Master Data Ops": ["master", "mifid", "mifir", "data", "emt", "attribute"],
}


def assign_role_owner(filename: str) -> str:
    name = filename.lower()
    for role, keywords in ROLE_KEYWORDS.items():
        if any(k in name for k in keywords):
            return role
    return random.choice(ROLE_OWNERS)


def random_recent_date() -> str:
    days_ago = random.randint(5, 400)
    d = datetime.date.today() - datetime.timedelta(days=days_ago)
    return d.isoformat()


def read_pdf(path: str) -> str:
    """Read any file as a PDF. Raises if it is not a real PDF."""
    text_parts = []
    with fitz.open(path) as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts)


def read_docx(path: str) -> str:
    doc = DocxDocument(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def read_xlsx(path: str) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows_text = []
    for ws in wb.worksheets:
        rows_text.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows_text.append(" | ".join(cells))
    wb.close()
    return "\n".join(rows_text)


def extract_text(path: str) -> str:
    """
    The dirty-data hack: try PDF first regardless of extension, because many
    files are PDFs wearing a .xlsx/.docx costume. Only if PyMuPDF fails do we
    fall back to genuine Office parsers.
    """
    try:
        text = read_pdf(path)
        if text.strip():
            return text
    except Exception:
        pass

    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".docx":
            return read_docx(path)
        if ext in (".xlsx", ".xlsm"):
            return read_xlsx(path)
    except Exception:
        return ""
    return ""


def main():
    # Reset the store so re-running does not append duplicate chunks.
    if os.path.isdir(CHROMA_DIR):
        print(f"Resetting existing vector store at ./{CHROMA_DIR}")
        shutil.rmtree(CHROMA_DIR)

    splitter = RecursiveCharacterTextSplitter(chunk_size=500, chunk_overlap=50)
    all_docs = []

    for filename in sorted(os.listdir(DATA_DIR)):
        path = os.path.join(DATA_DIR, filename)
        if not os.path.isfile(path) or filename.startswith("."):
            continue

        text = extract_text(path)
        if not text.strip():
            print(f"SKIP  (unreadable): {filename}")
            continue

        role_owner = assign_role_owner(filename)
        last_updated = random_recent_date()

        chunks = splitter.split_text(text)
        for chunk in chunks:
            all_docs.append(
                Document(
                    page_content=chunk,
                    metadata={
                        "source_file": filename,
                        "role_owner": role_owner,
                        "last_updated": last_updated,
                    },
                )
            )
        print(f"OK    {filename}: {len(chunks)} chunks  ->  {role_owner}")

    print(f"\nTotal chunks: {len(all_docs)}")
    if not all_docs:
        print("No documents to index. Aborting.")
        return

    embeddings = HuggingFaceEmbeddings(model_name=EMBED_MODEL)
    Chroma.from_documents(
        documents=all_docs,
        embedding=embeddings,
        collection_name=COLLECTION_NAME,
        persist_directory=CHROMA_DIR,
    )
    print(f"Saved vector store to ./{CHROMA_DIR}")


if __name__ == "__main__":
    main()
